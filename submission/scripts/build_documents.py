"""Create the research manuscript and technical project report from measured evidence."""
import json
import csv
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parents[1]
rows=list(csv.DictReader((OUT/'results/metrics.csv').open()))
for r in rows:
    for k in ['accuracy','precision','recall','f1','roc_auc','ap','threshold','train_seconds']:r[k]=float(r[k])
audit=json.loads((OUT/'results/dataset_audit.json').read_text());selection=json.loads((OUT/'results/selection.json').read_text())
analysis=json.loads((OUT/'results/analysis.json').read_text());boot=json.loads((OUT/'results/bootstrap.json').read_text())
manifest=json.loads((OUT/'results/run_manifest.json').read_text());parity=json.loads((OUT/'results/browser_parity.json').read_text())
names=['LogisticRegression','RandomForest','XGBoost','SVM','GradientBoosting','MLP','ResidualMLP','CNN1D']
labels=['Logistic Regression','Random Forest','XGBoost','SVM','Histogram GB','MLP','Residual MLP','1D CNN']
REFS=[
('[1] Project dataset','Health Insurance Fraud Claims.xlsx. Repository workbook, 4,500 records and 19 columns. Source origin, collection protocol and currency are undocumented. SHA-256 '+audit['source_sha256']+'.'),
('[2] Scikit learn','Average precision score. https://scikit-learn.org/stable/modules/generated/sklearn.metrics.average_precision_score.html Accessed 20 October 2026.'),
('[3] Scikit learn','Metrics and scoring. https://scikit-learn.org/stable/modules/model_evaluation.html Accessed 20 October 2026.'),
('[4] Scikit learn','Common pitfalls and recommended practices. https://scikit-learn.org/stable/common_pitfalls.html Accessed 20 October 2026.'),
('[5] PyTorch','BCEWithLogitsLoss. https://docs.pytorch.org/docs/stable/generated/torch.nn.BCEWithLogitsLoss.html Accessed 20 October 2026.'),
('[6] Gorishniy and colleagues','Y. Gorishniy, I. Rubachev, V. Khrulkov and A. Babenko. Revisiting Deep Learning Models for Tabular Data. NeurIPS 34, 2021. https://proceedings.neurips.cc/paper_files/paper/2021/hash/9d86d83f925f2149e9edb0ac3b49229c-Abstract.html'),
('[7] Experiment artifacts','submission/results/metrics.csv, predictions.csv, dataset_audit.json, selection.json, browser_parity.json and run_manifest.json. Reproducible internal experiment, 20 October 2026.')]

def get(s,n,policy='validation_f1',split='test'):
    return next(r for r in rows if r['scenario']==s and r['model']==n and r['policy']==policy and r['split']==split)
def fmt(v):return f'{float(v):.4f}'
def pct(v):return f'{100*float(v):.2f}%'
def p(doc,text,style=None):
    para=doc.add_paragraph(text,style);para.paragraph_format.widow_control=True;return para
def h(doc,text,level=1):return doc.add_heading(text,level)
def page(doc,title):
    q=h(doc,title);q.paragraph_format.page_break_before=True
def caption(doc,text):
    q=p(doc,text);q.paragraph_format.space_after=Pt(8)
    for r in q.runs:r.font.size=Pt(9);r.italic=True
def fig(doc,name,width=6.5):doc.add_picture(str(OUT/'figures'/name),width=Inches(width))
def base(kind):
    d=Document();s=d.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.70);s.left_margin=s.right_margin=Inches(.78)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=d.styles[name];st.font.name='Times New Roman' if kind=='paper' else 'Calibri';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(11);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.12
        fonts=st.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is not None:
            for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                fonts.attrib.pop(qn('w:'+k),None)
    # Remove decorative borders and tracking inherited from the packaged default template.
    for st in d.styles:
        for el in list(st.element.iter()):
            if el.tag in [qn('w:pBdr'),qn('w:spacing')] and el.getparent() is not None and el.getparent().tag==qn('w:rPr'):
                el.getparent().remove(el)
            elif el.tag==qn('w:pBdr') and el.getparent() is not None:el.getparent().remove(el)
    d.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    d.styles['Subtitle'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    d.styles['Subtitle'].font.italic=False
    d.styles['Title'].font.size=Pt(21);d.styles['Title'].font.bold=True
    d.styles['Heading 1'].font.size=Pt(15);d.styles['Heading 1'].paragraph_format.space_before=Pt(10)
    d.styles['Heading 2'].font.size=Pt(12);d.styles['Heading 2'].font.bold=True
    d.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.append(fld)
    d.core_properties.title='Medical Insurance Claim Fraud Detection'
    d.core_properties.subject='Approach 1 and Approach 2'
    d.core_properties.author='Project team'
    return d
def table(doc,headers,data,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=str(v)
    trpr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trpr.append(rep)
    for row in data:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'183B4A' if ri==0 else ('F1F5F6' if ri%2==0 else 'FFFFFF'));tc.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');margins.append(el)
            tc.append(margins)
            for q in c.paragraphs:
                q.paragraph_format.space_after=Pt(2);q.paragraph_format.space_before=Pt(2);q.paragraph_format.line_spacing=1.05
                q.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in q.runs:run.font.size=Pt(9.2);run.bold=ri==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
    p(doc,'').paragraph_format.space_after=Pt(0)
    return t
def results_table(doc,scenario='intake'):
    data=[]
    for n,label in zip(names,labels):
        r=get(scenario,n);data.append([label,fmt(r['ap']),pct(r['precision']),pct(r['recall']),fmt(r['f1']),f"{r['fp']} / {r['fn']}"])
    return table(doc,['Model','AP','Precision','Recall','F1','FP / FN'],data,[1.65,.75,.95,.95,.75,.65])
def sensitivity_table(doc):
    table(doc,['Model','Original AP','Intake AP','No amount AP','Temporal AP'],[[label]+[fmt(get(s,n)['ap']) for s in ['baseline','intake','no_amount','temporal']] for n,label in zip(names,labels)],[1.65,1,1,1,1.05])
def formulas(doc):
    p(doc,'Let TP and TN denote correctly flagged fraud and correctly unflagged legitimate claims. FP denotes a legitimate claim flagged for review; FN denotes a fraud claim below the decision threshold.')
    table(doc,['Measure','Definition'],[
        ['Accuracy','(TP + TN) / (TP + TN + FP + FN)'],['Precision','TP / (TP + FP)'],['Recall','TP / (TP + FN)'],['F1','2 TP / (2 TP + FP + FN)'],
        ['Average precision','Sum over score groups of (R_k - R_(k-1)) P_k'],['ROC AUC','Probability that a fraud score exceeds a legitimate score, with half credit for a tie']],[1.35,5.3])
    p(doc,'P_k and R_k are precision and recall after including a tied-score group in descending score order. AP uses recall increments as weights [2]. ROC AUC measures ranking across the two classes [3]. Zero-positive precision is defined as zero. Every score is uncalibrated.')
def training_text(doc):
    p(doc,'All neural models use weighted binary cross entropy with logits, with positive weight equal to the training negative-to-positive ratio (15.6667 for the random split) [5]. AdamW uses learning rate 0.001 and weight decay 0.0001. Batch size is 128. The maximum is 100 epochs, with patience 15 and checkpoint selection by validation AP. Seed 42, two CPU threads and deterministic PyTorch operations control the run. The saved best checkpoint, rather than the final epoch, supplies evaluation scores.')
def references(doc):
    for title,text in REFS:
        q=p(doc,title+'. '+text);q.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for r in q.runs:r.font.size=Pt(9)

def paper():
    d=base('paper');p(d,'Classical and Neural Models for Medical Insurance Claim Fraud Detection','Title')
    p(d,'A comparative study of Approach 1 and Approach 2', 'Subtitle')
    p(d,'20 October 2026')
    h(d,'Abstract')
    p(d,'This study compares five classical classifiers and three neural architectures for binary fraud detection on a repository workbook of 4,500 insurance claims, including 270 fraud labels. Approach 1 evaluates Logistic Regression, Random Forest, XGBoost, SVM and histogram Gradient Boosting. Approach 2 evaluates MLP, Residual MLP and 1D CNN. A shared stratified 70/15/15 split supports direct comparison. The protocol excludes post-adjudication status and identifiers, fits preprocessing only on training data and chooses decision thresholds using validation F1. An intake variant removes an unexplained Cluster field. Intake Random Forest achieves test average precision 1.0000 and F1 0.9873, with 39 true positives, one false negative and no false positives. However, removing both Cluster and claim amount reduces average precision to 0.2066 for that model. A chronological evaluation and a model-testing website complete the comparison. The evidence supports an internal benchmark demonstration and identifies substantial dependence on amount-related patterns. External validity remains unresolved.')
    p(d,'Keywords: insurance claims, fraud classification, imbalanced data, classical machine learning, deep learning, average precision.')
    h(d,'1 Introduction')
    p(d,'A fraud classifier can prioritize claims for human review by assigning scores to structured intake information. The useful comparison is whether a model separates rare fraud labels from legitimate claims while keeping the review workload manageable. Accuracy alone can obscure that distinction: predicting every claim as legitimate gives 94.00% accuracy and zero fraud recall in this workbook.')
    p(d,'The study addresses two questions within the existing project: how five classical methods compare with three neural architectures on identical partitions, and how sensitive the results are to uncertain dataset fields. It also connects the fitted models to a reproducible website demonstration. The contribution is a controlled evaluation and a functioning demonstration, rather than a new algorithm or evidence of performance on independent insurance operations.')
    page(d,'2 Dataset and experimental design')
    p(d,'The workbook contains 19 columns, no empty cells and no duplicate complete rows [1]. Claim dates span 9 July 2022 to 8 July 2024. The identifiers ClaimID, PatientID and ProviderID are unique in every row. DiagnosisCode and ProcedureCode each have 4,495 unique values. Such fields provide little reusable signal in a small benchmark and are excluded along with ProviderLocation.')
    table(d,['Partition','Claims','Fraud','Legitimate'],[['Training',3150,189,2961],['Validation',675,41,634],['Test',675,40,635]],[1.65,1.35,1.35,1.35])
    p(d,'ClaimLegitimacy maps Fraud to 1 and Legitimate to 0. Two stratified splits with seed 42 recreate the stored project partitions. Derived date features are year, month and weekday. Training-only median imputation and standard scaling process numeric inputs; mode imputation and one-hot encoding process six categorical fields. No imputation is needed for the current complete workbook. This fitting boundary follows standard leakage controls [4].')
    h(d,'2 1 Feature contracts',2)
    p(d,'The original comparison retains seven numeric and six categorical fields, yielding 29 encoded features. ClaimStatus is excluded because its adjudication timing could expose information unavailable at claim intake. The intake variant removes Cluster, yielding 28 features. Cluster 1 contains 267 of the 270 fraud labels, but its construction and availability are undocumented. The amount sensitivity variant also removes ClaimAmount, yielding 27 features. All three variants use the same random row partitions.')
    h(d,'2 2 Chronological assessment',2)
    p(d,'A separate intake experiment sorts by claim date and preserves complete dates at the boundaries. Training ends on 21 November 2023 (3,152 claims, 199 fraud cases). Validation runs from 22 November 2023 to 10 March 2024 (677 claims, 34 fraud cases). The test partition begins on 11 March 2024 (671 claims, 37 fraud cases). Preprocessing, neural checkpoint selection and thresholds are refitted within this chronological experiment. No calendar date appears in more than one partition.')
    p(d,'The random test split has appeared in earlier project analyses. Consequently, the study is a retrospective internal comparison, not a newly locked external assessment. The sensitivity experiments do not change the selected model or threshold after examining their test outcomes.')
    page(d,'3 Models and evaluation')
    table(d,['Type','Algorithm','Configuration'],[
        ['1','Logistic Regression','Balanced weights, lbfgs, maximum 2,000 iterations'],['1','Random Forest','300 trees, balanced weights'],['1','XGBoost','100 trees, depth 6, learning rate 0.3, training-derived class weight'],['1','SVM','RBF kernel, balanced weights, sigmoid of decision margin'],['1','Histogram GB','100 iterations, balanced weights'],
        ['2','MLP','Input, 64, 32, 1 units with ReLU and 0.2 dropout'],['2','Residual MLP','64-unit projection, two 64-unit residual blocks, scalar output'],['2','1D CNN','Two kernel-3 convolutions with 8 and 16 channels, flatten, 32-unit dense layer, scalar output']],[.6,1.4,4.65])
    training_text(d)
    p(d,'The residual model is a compact project implementation with additive skip connections. Work on tabular deep learning motivates comparing strong dense and residual baselines [6], but this model does not reproduce that paper. The CNN follows the fixed encoded feature order; neighboring columns have no established spatial meaning. It is therefore an experimental comparator.')
    p(d,'For each fitted model, the validation threshold maximizes F1 over unique validation scores. Equal F1 chooses the larger threshold. Primary model ranking uses validation AP rounded to 12 decimals, then validation F1, then the declared model order. The 0.5 policy is also retained to expose the effect of threshold choice. The test partition supplies final descriptive metrics only.')
    page(d,'4 Evaluation measures')
    formulas(d)
    p(d,'The project previously labelled average_precision_score output as PR AUC. This study calls the quantity average precision (AP) because it uses step weighting rather than trapezoidal interpolation [2]. The distinction matters when curve shapes differ. AP and ROC AUC use scores, while accuracy, precision, recall and F1 depend on a decision threshold.')
    h(d,'4 1 Independent verification',2)
    p(d,'The evaluation stores each validation and test label, source row and predicted score. Independent confusion-count equations reproduce the threshold metrics. A grouped-score implementation reproduces AP; a rank-sum calculation reproduces ROC AUC, including ties. These checks cover 128 metric rows across 32 fitted model and experiment combinations [7].')
    p(d,'For intake Random Forest at threshold 0.4167, TP = 39, TN = 635, FP = 0 and FN = 1. Accuracy is 674/675 = 99.8519%. Precision is 39/39 = 100.00%. Recall is 39/40 = 97.50%. F1 is 78/79 = 0.987342. Perfect AP can coexist with one false negative because ranking and a fixed operating threshold measure different properties.')
    fig(d,'confusion_matrix.png',3.5)
    caption(d,'Figure 1  Intake Random Forest test confusion matrix at the validation-selected threshold.')
    page(d,'5 Comparative results')
    p(d,'Table 2 reports the intake experiment with Cluster removed. Every row uses that model’s own validation-selected threshold. The models share the same 675 test claims and 40 fraud cases. Random Forest wins the specified validation tie-break and serves as the default demonstration model. Test results alone would not justify changing that selection.')
    results_table(d)
    caption(d,'Table 2  Intake test results. AP and F1 are fractions. Precision and recall are percentages. FP and FN are counts.')
    p(d,'XGBoost also achieves AP 1.0000 and has one false positive with no false negatives. Its F1 of 0.9877 marginally exceeds Random Forest’s 0.9873 at the selected thresholds, but the difference is one error of a different type. The best neural AP is 0.9667 for 1D CNN. Its selected threshold produces recall 80.00% and F1 0.8767. Neither architecture complexity nor AP alone guarantees a better operating point.')
    fig(d,'ap_comparison.png',6.5)
    caption(d,'Figure 2  Test AP for the original 29-feature comparison and the intake variant without Cluster.')
    page(d,'6 Sensitivity and uncertainty')
    sensitivity_table(d)
    caption(d,'Table 3  Test AP under four experiments. Each chronological model refits on its own earlier training period.')
    p(d,'Removing Cluster preserves strong tree performance but reduces AP for several linear and neural models. Removing claim amount as well produces AP values between 0.1869 and 0.2404 across all eight methods. Intake Random Forest falls from 1.0000 to 0.2066. All fraud amounts lie between 7,465.07 and 9,978.43 source units, although legitimate amounts overlap this range. The result indicates a strong amount-dependent pattern, rather than proof that amount alone determines fraud.')
    p(d,'The chronological test yields AP and F1 of 1.0000 for Random Forest, XGBoost and histogram Gradient Boosting. This internal temporal result suggests stable patterns within the workbook. It does not resolve unknown data provenance or establish transportability to another insurer or collection process.')
    p(d,'A stratified bootstrap with 1,000 replicates and seed 2026 keeps the fitted intake Random Forest and threshold fixed. The 95% percentile intervals are F1 0.9610 to 1.0000 and recall 0.9250 to 1.0000. The AP interval degenerates at 1.0000 because the observed scores perfectly rank the classes. These conditional intervals omit training variability, threshold-selection uncertainty and external distribution shift. They are not evidence that population AP equals one.')
    h(d,'6 1 Improvement priorities',2)
    p(d,'Removing the unexplained Cluster field improves the suitability of the input contract for a claim-intake demonstration. Additional priority work is to document the label and Cluster creation process, obtain independently labelled claims, evaluate across insurers and time, and calibrate scores on a separate partition. Thresholds should ultimately reflect measured review capacity and the relative costs of missed fraud and false alerts. The present study has neither those costs nor an external validation cohort.')
    page(d,'7 Demonstration and conclusions')
    p(d,'Claim Fraud Lab exposes both approaches through a shared form. A user selects the original or intake feature contract, chooses a model and threshold policy, and tests a claim or compares all eight algorithms. The browser evaluates exported fitted parameters locally. A same-origin FastAPI application provides an independent Python prediction endpoint. Inputs are validated, predictions display their thresholds, and results can be exported as JSON. The website labels all outputs as uncalibrated scores. The private demonstration is available at https://claim-fraud-lab-ravi-2026.clearbell16.chatgpt.site.')
    p(d,'Verification covers 1,350 raw-claim preprocessing comparisons and 10,800 model-score comparisons across the two website feature contracts. Preprocessing agrees exactly after float32 conversion. The maximum score difference is below 0.000001 and there are no threshold decision mismatches. The API passes 32 prediction requests plus invalid-input checks. This establishes implementation consistency on the tested cases, not predictive validity outside the workbook.')
    h(d,'7 1 Conclusion',2)
    p(d,'The completed comparison contains five classical methods and three neural methods, with shared evaluation, feature sensitivity tests and a functioning model interface. Tree models dominate this internal benchmark. Removing Cluster preserves a strong default Random Forest result, while amount ablation exposes a major dependency in the dataset. The appropriate submission claim is a reproducible comparative prototype with explicit limits, supported by saved predictions and measured results.')
    h(d,'References')
    references(d)
    path=OUT/'deliverables/Research_Paper.docx';d.save(path);return path

def report():
    d=base('report');p(d,'Medical Insurance Claim Fraud Detection','Title');p(d,'Technical Project Report for Approach 1 and Approach 2','Subtitle');p(d,'20 October 2026')
    h(d,'Project outcome')
    p(d,'The project compares five classical classifiers in Approach 1 with three neural classifiers in Approach 2. It includes a reproducible training and evaluation pipeline, 32 fitted model and experiment combinations, an interactive website, a research manuscript and a presentation. All result tables derive from the experiment artifacts generated on 20 October 2026.')
    p(d,'The default website configuration uses Random Forest without Cluster. On the random test split, it flags 39 of 40 fraud cases and no legitimate cases at threshold 0.4167. Its AP is 1.0000 and F1 is 0.9873. Removing claim amount as well reduces its AP to 0.2066. This dependency is central to interpreting the project.')
    table(d,['Component','Delivered scope'],[['Approach 1','Logistic Regression, Random Forest, XGBoost, SVM, histogram Gradient Boosting'],['Approach 2','MLP, Residual MLP and 1D CNN'],['Data','Original workbook, shared partitions, feature contracts and row manifests'],['Evaluation','Default and validation thresholds, amount ablation, temporal assessment and uncertainty'],['Demonstration','Browser inference plus local FastAPI endpoint'],['Submission files','Research paper, technical report and editable PowerPoint']],[1.45,5.2])
    h(d,'Reading guide')
    p(d,'The following chapters describe the problem, dataset, preprocessing, algorithms, evaluation, results, model sensitivity, website and reproduction steps. Appendices provide the artifact map, runtime versions and references. The report treats model scores as support for human review. It does not use the score to approve or reject insurance coverage.')
    page(d,'1 Problem and system scope')
    p(d,'The task is supervised binary classification of structured insurance claim records. A score ranks the tendency of a claim to resemble the workbook’s Fraud class. The demonstration converts that score to a review flag only after applying the displayed threshold. The intended audience is a project evaluator assessing methodology, measured comparison and implementation.')
    h(d,'1 1 Functional requirements',2)
    for text in ['Train and evaluate every existing classical algorithm using reproducible paths and a shared dataset split.','Extend the existing MLP with two supervised neural architectures, keeping Approach 2 at three algorithms.','Fit preprocessing only on training rows and retain validation and test predictions for audit.','Compare the original input contract with an intake contract that excludes an undocumented Cluster field.','Provide model selection, claim entry, navigation between approaches and comparison of model outputs in one website.','Provide consistent figures and results in a research manuscript, report and presentation.']:p(d,text,'List Bullet')
    h(d,'1 2 Implemented processing flow',2)
    p(d,'The offline flow reads the workbook, audits its structure, creates source-row partitions, fits the selected preprocessor, trains the eight algorithms, selects checkpoints and thresholds on validation data, and writes test scores and metrics. Export then converts the fitted parameters and preprocessing state into browser-readable JSON. The website applies the same transformations and model functions to user input.')
    p(d,'The optional local service loads the saved Python models and exposes /api/health and /api/predict. Static files and the API share http://127.0.0.1:8080. The hosted static website evaluates models in the visitor’s browser and therefore does not depend on a Python server being online.')
    page(d,'2 Dataset audit')
    p(d,'The supplied workbook contains 4,500 rows and 19 columns, of which 270 labels are Fraud and 4,230 are Legitimate [1]. The repository-root copy and Approach 1 raw-data copy have identical bytes. There are no missing cells or duplicate complete rows. The source hash is recorded in dataset_audit.json and the run manifest.')
    table(d,['Field group','Columns and handling'],[['Identifiers','ClaimID, PatientID and ProviderID are unique and excluded.'],['Claim details','ClaimAmount is retained except in the amount ablation. ClaimDate supplies year, month and weekday.'],['Near unique attributes','DiagnosisCode, ProcedureCode and ProviderLocation are excluded.'],['Patient attributes','Age, income, gender, marital status and employment status are retained.'],['Categorical intake fields','ProviderSpecialty, ClaimType and ClaimSubmissionMethod are retained.'],['Post adjudication attribute','ClaimStatus is excluded from every model.'],['Undocumented feature','Cluster is retained only for the original benchmark.'],['Target','ClaimLegitimacy maps Fraud to 1 and Legitimate to 0.']],[1.65,5])
    h(d,'2 1 Evidence requiring caution',2)
    table(d,['Cluster','Fraud','Legitimate'],[[0,0,1144],[1,267,833],[2,0,1104],[3,3,1149]],[1.65,2,2])
    p(d,'Cluster 1 contains 98.89% of the fraud records. All fraud claim amounts exceed 7,465.06 source units. The mean fraud amount is 8,733.57, compared with 4,776.80 for legitimate claims. Legitimate amounts overlap the fraud range. Dataset origin, currency, label adjudication and field timing are undocumented, so these patterns may reflect the collection or construction process.')
    page(d,'3 Preprocessing and partition controls')
    table(d,['Experiment','Encoded fields','Removed from original contract'],[['Original',29,'None beyond common exclusions'],['Intake',28,'Cluster'],['Amount sensitivity',27,'Cluster and ClaimAmount'],['Chronological intake',28,'Cluster']],[1.7,1.2,3.75])
    p(d,'Common exclusions are ClaimStatus, ClaimID, PatientID, ProviderID, DiagnosisCode, ProcedureCode and ProviderLocation. ClaimDate itself is replaced by year, month and weekday. The target never enters the feature transformer. Random partition indices match the original Approach 1 arrays for every split after preprocessing, allowing a direct comparison.')
    p(d,'Numeric columns pass through a training-fitted median imputer and StandardScaler. Categorical columns pass through a training-fitted most-frequent imputer and OneHotEncoder. The current workbook requires no filling. The training encoder controls column order and unseen-category handling. The website accepts only the known categorical choices to avoid silent unknown-field behavior. These operations fit on the training partition and transform the remaining partitions without refitting [4].')
    table(d,['Split','Claims','Fraud','Fraud share'],[['Random training',3150,189,'6.00%'],['Random validation',675,41,'6.07%'],['Random test',675,40,'5.93%'],['Chronological training',3152,199,'6.31%'],['Chronological validation',677,34,'5.02%'],['Chronological test',671,37,'5.51%']],[2,1.35,1.35,1.95])
    p(d,'Random partitioning uses 70% training, 15% validation and 15% test with stratification and seed 42. Chronological training ends on 21 November 2023 and validation ends on 10 March 2024. Keeping date groups intact changes exact chronological counts. Split manifests record workbook row numbers and claim identifiers to make partition membership auditable.')
    page(d,'4 Approach 1 classical algorithms')
    table(d,['Model','Implementation and training settings'],[['Logistic Regression','Linear log odds, balanced weights, lbfgs, max_iter 2000, seed 42.'],['Random Forest','300 trees, balanced class weights, seed 42 and two parallel jobs.'],['XGBoost','100 trees, maximum depth 6, learning rate 0.3, log loss objective, training negative-to-positive class ratio.'],['SVM','RBF kernel, balanced class weights and default C = 1. The sigmoid of the decision margin provides a bounded uncalibrated score.'],['Histogram Gradient Boosting','100 boosting iterations, balanced class weights, seed 42. Registry name GradientBoosting refers to HistGradientBoostingClassifier.']],[1.7,4.95])
    p(d,'The linear model provides a compact baseline. Random Forest averages decisions from separately grown trees. XGBoost and histogram Gradient Boosting learn successive trees to correct earlier errors. SVM separates the encoded feature vectors using an RBF kernel. The same encoded arrays feed every classical and neural model within each experiment.')
    h(d,'4 1 Existing artifact verification',2)
    p(d,'The original five classical models and original MLP checkpoint load and produce evaluation scores. Their metrics appear in legacy_artifact_audit.json. The original XGBoost artifact issues a serialization-version warning but evaluates successfully. Fresh models are saved separately under submission/models so that historical artifacts remain available for comparison.')
    p(d,'The historical SVM uses separate prediction and probability paths. At probability threshold 0.5 its F1 differs from the historical .predict result. The fresh pipeline uses one sigmoid-margin score for all threshold decisions, avoiding that ambiguity. The original comparison’s SVM ranking metrics remain reproducible, but its threshold-dependent numbers must be read with their specific decision rule.')
    p(d,'No grid search chooses hyperparameters from the test data. The configuration table defines a bounded model comparison, rather than a claim of globally optimal hyperparameters. Weighted training addresses the observed class imbalance without generating synthetic claim rows.')
    page(d,'5 Approach 2 neural algorithms')
    table(d,['Model','Architecture'],[['MLP','Input of 27, 28 or 29 features, dense 64, ReLU, dropout 0.2, dense 32, ReLU, dropout 0.2, scalar logit.'],['Residual MLP','Input projection to 64, ReLU, two residual blocks. Each block uses dense 64, ReLU, dropout 0.2 and dense 64, adds its input and applies ReLU. A linear output yields a logit.'],['1D CNN','Treat the fixed feature vector as one channel. Conv1d 1 to 8 and 8 to 16, both kernel 3 and padding 1, each followed by ReLU. Flatten, dropout 0.2, dense 32, ReLU and scalar logit.']],[1.5,5.15])
    training_text(d)
    table(d,['Intake model','Best epoch','Epochs run','Parameters'],[[label,json.loads((OUT/'results'/f'intake_{n}_training.json').read_text())['best_epoch'],json.loads((OUT/'results'/f'intake_{n}_training.json').read_text())['epochs'],json.loads((OUT/'results'/f'intake_{n}_training.json').read_text())['parameters']] for n,label in zip(names[5:],labels[5:])],[2,.95,1.3,1.6])
    fig(d,'neural_training.png',6.2)
    caption(d,'Figure 1  Intake validation AP by epoch. Each architecture restores its own best validation checkpoint.')
    p(d,'The residual connections provide an additional dense comparator. The CNN uses a feature ordering fixed by preprocessing, although categorical and numeric columns have no physical spatial structure. Its outcome is specific to this encoding. The comparison does not assume that deep learning must outperform tree ensembles on tabular data [6].')
    page(d,'6 Evaluation and threshold selection')
    formulas(d)
    p(d,'A fixed all-legitimate classifier obtains 94.00% accuracy on the full workbook and 94.0741% on the random test set, with zero fraud recall and F1. AP for a constant score equals prevalence. This provides a more informative reference than accuracy alone. Random-test prevalence is 40/675 = 0.0593.')
    p(d,'Each model selects the threshold that maximizes validation F1. Ties choose the larger threshold. Ranking for the default intake model uses validation AP rounded to 12 decimals, validation F1 and finally the declared model order. The rule selects Random Forest. Test scores do not select a model or threshold. The complete results retain both the 0.5 and validation-F1 policies.')
    h(d,'6 1 Worked numerical check',2)
    p(d,'For intake Random Forest, the validation threshold is 0.4166667. On the test split, TN = 635, FP = 0, FN = 1 and TP = 39. Accuracy = (635 + 39)/675 = 99.8519%. Precision = 39/(39 + 0) = 100.00%. Recall = 39/(39 + 1) = 97.50%. F1 = 78/(78 + 0 + 1) = 0.987342. The same counts and equations appear in the saved evaluation evidence.')
    p(d,'The verification script recomputes threshold metrics from counts and independently computes AP from tied score groups and ROC AUC from ranks. All 128 rows reconcile. A saved CSV prediction is therefore the traceable basis of every reported result, rather than a copied plot label.')
    page(d,'7 Main comparison results')
    results_table(d)
    caption(d,'Table 7  Intake experiment test results at validation-selected thresholds. Each algorithm uses the same 675 test rows.')
    p(d,'Random Forest and XGBoost have perfect test ranking in the intake experiment. Random Forest misses one fraud claim, while XGBoost flags one legitimate claim. Histogram Gradient Boosting has the same count-level result as Random Forest, although its AP is lower. This illustrates why ranking quality and threshold quality should be reported together.')
    p(d,'The 1D CNN has the highest neural AP, 0.9667. Its threshold of 0.9691 produces one false positive and eight false negatives. MLP and Residual MLP each detect 37 fraud cases but generate 13 and 12 false positives respectively. The operating point determines whether ranking improvements translate into a useful review workload.')
    fig(d,'selected_curves.png',6.5)
    caption(d,'Figure 2  Intake Random Forest precision recall and ROC curves. The dashed PR reference marks test prevalence.')
    page(d,'8 Effects of threshold adjustment')
    table(d,['Model','F1 at 0.5','F1 at validation threshold','Chosen threshold'],[[label,fmt(get('intake',n,'default')['f1']),fmt(get('intake',n)['f1']),fmt(get('intake',n)['threshold'])] for n,label in zip(names,labels)],[1.8,1.2,2,1.65])
    p(d,'Validation-based threshold selection can improve an operating point, but the gain need not transfer to the test partition. The table shows both policies without selecting between them from test outcomes. Because validation contains only 41 fraud cases, a threshold can be sensitive to individual observations. A separate calibration or policy-tuning cohort would provide a stronger basis for production decisions.')
    fig(d,'confusion_matrix.png',4.6)
    caption(d,'Figure 3  Default demonstration model at its frozen validation-selected threshold.')
    h(d,'8 1 Score interpretation',2)
    p(d,'All outputs are uncalibrated model scores. Class weighting changes the training objective, and the SVM score is a transformed margin. Thus a score of 0.80 does not establish an 80% probability of fraud. The website displays this distinction next to the result. Scores from different algorithms should not be averaged as if they shared a probability scale.')
    page(d,'9 Feature sensitivity and error analysis')
    sensitivity_table(d)
    caption(d,'Table 9  Test AP across the original, intake, amount-removed and chronological experiments.')
    fig(d,'sensitivity.png',6.2)
    caption(d,'Figure 4  Random Forest sensitivity to the selected feature contract and the chronological split.')
    p(d,'Removing Cluster does not remove the strong tree ranking. Removing both Cluster and ClaimAmount lowers AP across all algorithms to 0.1869–0.2404. Random Forest AP becomes 0.2066. Training impurity importance assigns about 44.0% to PatientIncome and 43.2% to ClaimAmount in the intake Random Forest. These importance scores describe fitted splits and are neither causal effects nor unbiased evidence of independent relevance.')
    err=analysis['errors'][0]
    p(d,f"The intake Random Forest has one false negative at workbook row {err['source_row']}, with claim amount {err['ClaimAmount']:,.2f} source units, age {err['PatientAge']} and claim type {err['ClaimType']}. Its score is {err['score']:.4f}, below the frozen threshold. The record belongs to the evaluation evidence; it does not justify changing the trained model after seeing the test result.")
    page(d,'10 Chronological results and uncertainty')
    results_table(d,'temporal')
    caption(d,'Table 10  Chronological intake experiment. Test contains 671 claims and 37 fraud labels.')
    p(d,'Random Forest, XGBoost and histogram Gradient Boosting each obtain AP and F1 of 1.0000 on the later test period. The neural models have lower F1, ranging from 0.6957 for MLP to 0.8108 for CNN. All preprocessing and thresholds come from the earlier partitions. These dates assess internal temporal separation while remaining within the same undocumented workbook.')
    h(d,'10 1 Conditional uncertainty',2)
    p(d,'The selected random-split intake Random Forest has a 95% stratified bootstrap interval of 0.9610–1.0000 for F1 and 0.9250–1.0000 for recall, using 1,000 resamples and seed 2026. The method holds the model and threshold fixed and resamples positive and negative test rows separately. The AP interval is 1.0000–1.0000 because all observed test fraud scores outrank all legitimate scores. This degeneracy does not mean certainty about external population performance.')
    p(d,'A single training seed and one split per experiment limit the comparison. The bootstrap does not include training variance, repeated threshold selection, changes in class prevalence or unknown shifts in label creation. Historical reuse of the random test split further limits strong generalization claims. More trustworthy improvement requires a documented independent dataset and repeated training within a prespecified protocol.')
    page(d,'11 Website and prediction interface')
    p(d,'Claim Fraud Lab provides four sections: project overview, model lab, measured results and method. The model lab switches between Approach 1 and Approach 2, selects one of the eight algorithms and selects either the 29-feature benchmark or 28-feature intake contract. Two examples from the workbook allow a quick demonstration. All model outputs include the decision threshold and feature count.')
    table(d,['Interface element','Behavior'],[['Feature contract','Intake omits Cluster. Original comparison requires Cluster. ClaimStatus and target are never model inputs.'],['Claim form','Checks required numbers, date validity, age range and known category choices.'],['Model inference','Evaluates exported fitted model parameters in the browser. It does not call an external prediction service.'],['Comparison','Shows all eight scores and indicates whether each exceeds its own threshold.'],['Results view','Switches among original, intake, amount sensitivity and chronological experiments.'],['Prediction export','Downloads the inputs, contract, threshold policy and model outputs as JSON.'],['Local API','POST /api/predict computes the same contract with Python artifacts. GET /api/health lists eight models.']],[1.5,5.15])
    p(d,'The local site serves from http://127.0.0.1:8080. The private hosted copy is available at https://claim-fraud-lab-ravi-2026.clearbell16.chatgpt.site and uses the same static files and fitted parameter JSON. Browser inference allows the hosted demonstration to work without uploading claim inputs or exposing the local Python service. The two rendering paths serve the same completed model scope.')
    h(d,'11 1 Implementation validation',2)
    p(d,f"Across 1,350 test-claim preprocessing comparisons, browser and Python feature vectors match exactly after float32 conversion. Across 10,800 model-score comparisons, the largest absolute difference is {max(r['max_absolute_error'] for r in parity):.8f}, with zero decision mismatches. API checks cover 32 valid prediction requests, four invalid-input cases and five static routes. Browser interaction checks cover example loading, model comparison and navigation.")
    page(d,'12 Reproduction and demonstration procedure')
    h(d,'12 1 Environment',2)
    p(d,'Use Python 3.12 and the package versions in submission/requirements-lock.txt. Create a virtual environment at the repository root, then install dependencies. The tested environment uses CPU PyTorch. A CPU wheel index may be used for torch. The document builder uses the separate bundled artifact runtime and is not required to launch the model demonstration.')
    for text in ['python -m venv .submission-venv','.submission-venv\\Scripts\\python.exe -m pip install -r submission\\requirements-lock.txt','.submission-venv\\Scripts\\python.exe submission\\scripts\\run_pipeline.py','.submission-venv\\Scripts\\python.exe submission\\scripts\\export_browser.py','.submission-venv\\Scripts\\python.exe submission\\scripts\\make_figures.py','.submission-venv\\Scripts\\python.exe submission\\scripts\\verify.py','node submission\\scripts\\check_browser.mjs','powershell -ExecutionPolicy Bypass -File submission\\START_DEMO.ps1']:
        q=p(d,text);q.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for r in q.runs:r.font.name='Consolas';r.font.size=Pt(9)
    h(d,'12 2 Demonstration sequence',2)
    for text in ['Open the site and state the two approaches and the dataset class imbalance.','Use the intake contract and default Random Forest. Load the legitimate example and test it.','Load the fraud example and compare all eight algorithms. Explain score and threshold differences.','Switch to Approach 2 and test each neural architecture.','Switch the feature contract to the original benchmark and show that Cluster becomes an input.','Open Results and compare the amount sensitivity and chronological experiments. End with the amount dependency and external-validation limitation.']:p(d,text,'List Number')
    page(d,'13 Artifact map and research limits')
    table(d,['Path under submission','Purpose'],[['scripts/run_pipeline.py','Shared data audit, training, validation, testing and selection'],['models/<experiment>/','Fitted preprocessors, five classical models, three neural checkpoints and contract'],['results/metrics.csv','128 metric records for partition and threshold policy combinations'],['results/predictions.csv','Auditable validation and test scores by workbook row'],['data/*_split_manifest.csv','Source-row partition membership'],['results/legacy_artifact_audit.json','Read-only evaluation of six historical trained artifacts'],['results/browser_parity.json','Browser and Python score consistency evidence'],['results/verification.json','Metric formulas and API checks'],['website/dist/','Portable static model laboratory and parameter JSON'],['deliverables/','Research manuscript, report and PowerPoint presentation']],[2.8,3.85])
    p(d,'The source workbook and reference documents remain intact. The new submission pipeline resolves paths from its own location, so it does not rely on the old project/ directory used in several historical scripts. The website package contains only Approach 1 and Approach 2 model functions and experiment results.')
    h(d,'13 1 Next research work',2)
    p(d,'Document data origin, informed field timing and the label-adjudication process. Determine whether Cluster or any numerical feature was constructed using target information. Evaluate new claims collected independently of this workbook. Add repeated training and uncertainty over model fitting. Calibrate scores using a dedicated cohort and set review thresholds using explicit workload and error costs. Assess demographic subgroup performance only with adequate sample sizes and an agreed evaluation purpose.')
    p(d,'The completed submission supports a reproducible model comparison and demonstration. The strong amount dependence, small positive test counts and unknown source provenance prevent a claim of production readiness for insurer decision making.')
    page(d,'References and runtime record')
    references(d)
    table(d,['Runtime component','Measured version'],[['Python',manifest['python']]]+[[k,v] for k,v in manifest['packages'].items()],[2.5,4.15])
    p(d,f"The final full training run took {manifest['seconds']:.2f} seconds on the local CPU environment. This is an observed run duration, not a hardware-independent speed claim. Random seed is 42. Bootstrap seed is 2026.")
    path=OUT/'deliverables/Project_Report.docx';d.save(path);return path

if __name__=='__main__':
    print(paper());print(report())
